"""Build committed fixtures with fixed seeds (P2-3, demo).

Outputs (all generated, reproducible):
  data/templates/synthetic_template.xlsx + synthetic_map.yaml
    A clearly labelled STAND-IN allocation template used until the
    organiser template is available. Contains merges, a validation, a
    formula, column widths and styled headers for the fidelity test.
"""
from pathlib import Path
import yaml
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
TDIR = ROOT / "data" / "templates"

FIELDS = ["tag", "cable", "pair", "cabinet", "cioc", "baseplate", "slot",
          "charm", "tb", "cabinet_name", "cioc_name", "baseplate_name",
          "io_type", "subtype", "group", "seg"]
HEADERS = ["Tag", "Cable", "Pair", "Cabinet", "CIOC", "Baseplate", "Slot",
           "CHARM", "Terminal Block", "Cabinet Name", "CIOC Name",
           "Baseplate Name", "IO Type", "Subtype", "Group", "Seg"]


def make_synthetic_template():
    TDIR.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Allocation"
    # Merged title row (must survive the fidelity test)
    ws.merge_cells("A1:P1")
    ws["A1"] = "CHARMSAT SYNTHETIC TEMPLATE - stand-in until the organiser template is available"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A1"].alignment = Alignment(horizontal="center")
    # Header row with style + widths
    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="1B4D3E", end_color="1B4D3E", fill_type="solid")
        ws.column_dimensions[get_column_letter(i)].width = max(10, len(h) + 2)
    # Data validation on IO Type column (M)
    dv = DataValidation(type="list", formula1='"AI,AO,DI,DO,RTD,TC"', allow_blank=True)
    dv.sqref = "M3:M10000"
    ws.add_data_validation(dv)
    # Summary block with a formula (must survive; writer never targets it)
    ws["R2"] = "Signals written"
    ws["S2"] = "=COUNTA(A3:A10000)"
    ws.freeze_panes = "A3"
    wb.save(TDIR / "synthetic_template.xlsx")

    tmap = {
        "template": "synthetic",
        "allocation_sheet": "Allocation",
        "header_row": 2,
        "first_data_row": 3,
        "columns": {f: get_column_letter(i + 1) for i, f in enumerate(FIELDS)},
        "summary_cells": {},
        "extra_sheets": True,
        "allow_formula_overwrite": [],
    }
    (TDIR / "synthetic_map.yaml").write_text(yaml.safe_dump(tmap, sort_keys=False))
    print(f"wrote {TDIR / 'synthetic_template.xlsx'} + synthetic_map.yaml")


if __name__ == "__main__":
    make_synthetic_template()
