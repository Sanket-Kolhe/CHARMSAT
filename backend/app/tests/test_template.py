"""P2-3: template writer - fidelity, refusal, XML fallback, roundtrip."""
import io
import shutil
import zipfile
from pathlib import Path

import pytest
import openpyxl

from app.engine.io.template_writer import (
    load_template_map, write_template, snapshot_workbook, allocation_hash,
    rows_from_allocations, TemplateWriteError, has_drawings,
)

TDIR = Path("data/templates")
TMAP = load_template_map(TDIR / "synthetic_map.yaml")


def _rows(n=30):
    return [{"tag": f"TT-{i}", "cable": "CBL-1", "pair": (i % 12) + 1,
             "cabinet": 1, "cioc": 1, "baseplate": (i // 12) + 1, "slot": (i % 12) + 1,
             "charm": "AI 4-20 mA HART", "tb": "Standard",
             "cabinet_name": "CAB-01", "cioc_name": "AR1CIOC100001",
             "baseplate_name": "001-CBP-1", "io_type": "AI", "subtype": "",
             "group": "G", "seg": "A"} for i in range(n)]


def test_write_preserves_template_structure(tmp_path):
    out = tmp_path / "filled.xlsx"
    before = snapshot_workbook(TDIR / "synthetic_template.xlsx")
    write_template(TDIR / "synthetic_template.xlsx", out, TMAP, _rows(),
                   extras={"TQ_LOG": [["id", "message"]], "AUDIT": [["hash", allocation_hash(_rows())]]})
    after = snapshot_workbook(out)
    # Structural snapshot identical except added extra sheets + grown data range
    assert after["sheets"] == before["sheets"] + ["TQ_LOG", "AUDIT"]
    alloc_b, alloc_a = before["per_sheet"]["Allocation"], after["per_sheet"]["Allocation"]
    for key in ("merges", "validations", "cond_formats", "widths", "formulas", "header_styles"):
        assert alloc_a[key] == alloc_b[key], key
    assert alloc_a["formulas"].get("S2") == "=COUNTA(A3:A10000)"
    # Values landed
    wb = openpyxl.load_workbook(out)
    ws = wb["Allocation"]
    assert ws["A3"].value == "TT-0" and ws["M3"].value == "AI"
    assert ws["S2"].value == "=COUNTA(A3:A10000)"


def test_refuses_formula_overwrite(tmp_path):
    # Craft a template whose data cell B3 holds a formula; the writer must
    # refuse rather than silently destroy it.
    src = tmp_path / "form.xlsx"
    shutil.copy(TDIR / "synthetic_template.xlsx", src)
    wb = openpyxl.load_workbook(src)
    wb["Allocation"]["B3"] = "=A3"
    wb.save(src)
    out = tmp_path / "x.xlsx"
    with pytest.raises(TemplateWriteError, match="formula cell"):
        write_template(src, out, TMAP, _rows(3))


def test_missing_sheet_raises(tmp_path):
    out = tmp_path / "x.xlsx"
    with pytest.raises(TemplateWriteError, match="not in template"):
        write_template(TDIR / "synthetic_template.xlsx", out,
                       dict(TMAP, allocation_sheet="Nope"), _rows(1))


def _drawing_template(tmp_path: Path) -> Path:
    """Template copy with an injected drawing part (simulates charts/images)."""
    src = tmp_path / "draw.xlsx"
    shutil.copy(TDIR / "synthetic_template.xlsx", src)
    drawing = ('<?xml version="1.0" encoding="UTF-8"?>'
               '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"/>')
    with zipfile.ZipFile(src, "a") as z:
        z.writestr("xl/drawings/drawing1.xml", drawing)
    return src


def test_xml_path_preserves_drawings(tmp_path):
    src = _drawing_template(tmp_path)
    assert has_drawings(src)
    out = tmp_path / "filled.xlsx"
    mode = write_template(src, out, TMAP, _rows(5))
    assert mode == "xml"
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert "xl/drawings/drawing1.xml" in names
        raw = z.read("xl/drawings/drawing1.xml").decode()
        assert "wsDr" in raw  # byte-preserved drawing part
    wb = openpyxl.load_workbook(out)
    assert wb["Allocation"]["A3"].value == "TT-0"


def test_allocation_hash_deterministic():
    assert allocation_hash(_rows()) == allocation_hash(_rows())
    assert allocation_hash(_rows(5)) != allocation_hash(_rows(6))
